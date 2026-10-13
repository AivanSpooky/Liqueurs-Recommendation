# -*- coding: utf-8 -*-
"""
ЛР 2. Меры близости и расстояния для гетерогенного датасета ликёров.

Части:
  ЛР 2.1 — частные меры на числовых признаках (>= 3 кандидатов),
           ассоциативные бинарные меры для bitter / wood_contact;
           экспертная оценка и выбор одной меры.
  ЛР 2.2 — мера близости по иерархической таксономии вкуса (Wu-Palmer).
  ЛР 2.3 — итоговая мера: взвешенная сумма частных мер по всем признакам
           объекта, кроме id и name.

Выход:
  out/matrices/*.csv   — матрицы смежности G_ij для каждой меры
  out/heatmaps/*.png   — тепловые карты
  out/report.txt       — числа для экспертной оценки
"""

import csv
import json
import os
import sys
import math

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
DATASET = os.path.join(HERE, "..", "1-dataset", "dataset.json")
OUT_M = os.path.join(HERE, "out", "matrices")
OUT_H = os.path.join(HERE, "out", "heatmaps")
OUT_R = os.path.join(HERE, "out", "report.txt")

os.makedirs(OUT_M, exist_ok=True)
os.makedirs(OUT_H, exist_ok=True)


# ------------------------------------------------------------------
# Загрузка данных
# ------------------------------------------------------------------
with open(DATASET, encoding="utf-8") as f:
    data = json.load(f)

recs = data["records"]
n = len(recs)
ids = [r["id"] for r in recs]
names = [r["name"] for r in recs]
labels = [f'{r["id"]} {r["name"][:18]}' for r in recs]


# ------------------------------------------------------------------
# Признаки
# ------------------------------------------------------------------
NUM_FIELDS = ["abv_pct", "sugar_gpl", "volume_l"]

X_raw = np.array([[r[f] for f in NUM_FIELDS] for r in recs], dtype=float)
x_min = X_raw.min(axis=0)
x_max = X_raw.max(axis=0)
# min-max нормировка: все оси в [0,1], единый масштаб,
# чтобы сахар (г/л) не перевешивал объём (л)
X = (X_raw - x_min) / (x_max - x_min)

B = np.array([[int(r["bitter"]), int(r["wood_contact"])] for r in recs])  # бинарные

# порядковая лингвистическая переменная по сахару (4 терма)
SUGAR_ORD = {"малосладкий": 0, "умеренный": 1, "сладкий": 2, "crème": 3}
sugar_ord = np.array([SUGAR_ORD[r["sugar_class"]] for r in recs], dtype=float)

countries = [r["country"] for r in recs]
producers = [r["producer"] for r in recs]
taste = [r["taste"] for r in recs]


# ------------------------------------------------------------------
# Утилиты
# ------------------------------------------------------------------
ALL_MAT = {}  # name -> (matrix, "sim" | "dist") — для xlsx и html


def save_csv(M, fname, kind=None):
    # разделитель ";" — чтобы русский Excel сразу разбивал по ячейкам
    path = os.path.join(OUT_M, fname)
    try:
        f = open(path, "w", newline="", encoding="utf-8-sig")
    except PermissionError:
        print(f"!! {path} занят (открыт в Excel?) — пропускаю запись")
        if kind:
            ALL_MAT[fname[:-4]] = (M, kind)
        return path
    with f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["id"] + ids)
        for i, row in enumerate(M):
            w.writerow([labels[i]] + [f"{v:.4f}".replace(".", ",") for v in row])
    if kind:
        ALL_MAT[fname[:-4]] = (M, kind)
    return path


def heatmap(M, fname, title, vmin=0.0, vmax=1.0, cmap="viridis"):
    fig, ax = plt.subplots(figsize=(13, 11))
    im = ax.imshow(M, cmap=cmap, vmin=vmin, vmax=vmax)
    ax.set_xticks(range(n))
    ax.set_yticks(range(n))
    ax.set_xticklabels(labels, rotation=90, fontsize=6)
    ax.set_yticklabels(labels, fontsize=6)
    ax.set_title(title, fontsize=11)
    fig.colorbar(im, ax=ax, shrink=0.8)
    fig.tight_layout()
    path = os.path.join(OUT_H, fname)
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def upper_tri(M):
    iu = np.triu_indices(n, 1)
    return M[iu]


def spearman(a, b):
    """Корреляция Спирмена двух векторов (ранги со средними для связей)."""
    def rank(v):
        order = np.argsort(v, kind="mergesort")
        r = np.empty(len(v))
        r[order] = np.arange(len(v))
        # средние ранги для связей
        sv = v[order]
        i = 0
        while i < len(v):
            j = i
            while j + 1 < len(v) and sv[j + 1] == sv[i]:
                j += 1
            if j > i:
                r[order[i:j + 1]] = (i + j) / 2.0
            i = j + 1
        return r
    ra, rb = rank(a), rank(b)
    ra = ra - ra.mean()
    rb = rb - rb.mean()
    denom = math.sqrt((ra ** 2).sum() * (rb ** 2).sum())
    return float((ra * rb).sum() / denom) if denom else 1.0


report = []


def rep(line=""):
    report.append(line)
    print(line)


# ==================================================================
# ЛР 2.1a. Числовые признаки: кандидатные меры расстояния
# (все работают на min-max нормированных координатах, область [0,1])
# ==================================================================
D = X.shape[1]
diff = np.abs(X[:, None, :] - X[None, :, :])          # |x_d - y_d|

d_euclid = np.sqrt((diff ** 2).sum(axis=2) / D)      # норм. евклидово (L2)
d_manh = diff.sum(axis=2) / D                        # норм. манхэттенское (L1)
d_mink3 = (diff ** 3).sum(axis=2) / D
d_mink3 = d_mink3 ** (1.0 / 3.0)                     # Минковский p=3 (к Чебышёву)
d_cheb = diff.max(axis=2)                            # чебышёвское (L∞)

norms = np.linalg.norm(X, axis=1)
cos_den = norms[:, None] * norms[None, :]
with np.errstate(divide="ignore", invalid="ignore"):
    cos_sim = np.where(cos_den > 0, (X @ X.T) / np.where(cos_den == 0, 1, cos_den), 1.0)
d_cosine = 1.0 - np.clip(cos_sim, -1.0, 1.0)         # косинусное расстояние

numeric_measures = {
    "euclidean": d_euclid,
    "manhattan": d_manh,
    "minkowski_p3": d_mink3,
    "chebyshev": d_cheb,
    "cosine": d_cosine,
}

for name_m, M in numeric_measures.items():
    save_csv(M, f"num_{name_m}_dist.csv", "dist")
    heatmap(M, f"num_{name_m}_dist.png",
            f"Числовые признаки — расстояние ({name_m})", cmap="magma")

# ==================================================================
# ЛР 2.1b. Бинарные признаки bitter, wood_contact:
# ассоциативные меры близости (Строганов, 2014)
# a = оба 1, b/c = рассогласования, d = оба 0
# ==================================================================
a = np.zeros((n, n)); b_ = np.zeros((n, n)); c_ = np.zeros((n, n)); d_ = np.zeros((n, n))
for i in range(n):
    for j in range(n):
        x, y = B[i], B[j]
        a[i, j] = np.sum((x == 1) & (y == 1))
        b_[i, j] = np.sum((x == 1) & (y == 0))
        c_[i, j] = np.sum((x == 0) & (y == 1))
        d_[i, j] = np.sum((x == 0) & (y == 0))


def safe_div(num, den, default=1.0):
    return np.where(den > 0, num / np.where(den == 0, 1, den), default)


binary_measures = {
    # Жаккар: учитывает только совпадения по 1; (0,0)-(0,0) -> 0/0 -> задаём 1
    "jaccard": safe_div(a, a + b_ + c_, default=1.0),
    # Дайс (Сёренсен): удвоенный вес совпадений по 1, d игнорируется
    "dice": safe_div(2 * a, 2 * a + b_ + c_, default=1.0),
    # Сокал–Миченер (simple matching): доля совпадающих позиций, включая 0-0
    "sokal_michener": (a + d_) / B.shape[1],
    # Роджерс–Танимото: как SMC, но рассогласование штрафуется вдвое
    "rogers_tanimoto": safe_div(a + d_, a + d_ + 2 * (b_ + c_)),
}

for name_m, M in binary_measures.items():
    save_csv(M, f"bin_{name_m}_sim.csv", "sim")
    heatmap(M, f"bin_{name_m}_sim.png",
            f"Бинарные признаки — близость ({name_m})")

# ==================================================================
# Номинальные и порядковые признаки (частные меры для агрегата)
# ==================================================================
def delta_measure(vals):
    M = np.zeros((n, n))
    for i in range(n):
        for j in range(n):
            M[i, j] = 1.0 if vals[i] == vals[j] else 0.0
    return M


sim_country = delta_measure(countries)
sim_producer = delta_measure(producers)
sim_sugar_class = 1.0 - np.abs(sugar_ord[:, None] - sugar_ord[None, :]) / 3.0

save_csv(sim_country, "nom_country_sim.csv", "sim")
save_csv(sim_producer, "nom_producer_sim.csv", "sim")
save_csv(sim_sugar_class, "ord_sugarclass_sim.csv", "sim")
heatmap(sim_country, "nom_country_sim.png", "Страна — близость (δ-функция)")
heatmap(sim_sugar_class, "ord_sugarclass_sim.png",
        "Класс сладости — близость (порядковая, 4 терма)")

# ==================================================================
# ЛР 2.2. Таксономия вкуса: мера У–Палмера
# sim(x,y) = 2*depth(LCA) / (depth(x)+depth(y)),
# листья на глубине 3 (root=0, L1=1, L2=2, L3=3).
# Эквивалентно взвешенному общему пути с весами рёбер w1=w2=w3=1.
# ==================================================================
def taste_sim(tx, ty):
    shared = (tx["l1"] == ty["l1"]) + (tx["l2"] == ty["l2"]) + (tx["l3"] == ty["l3"])
    if tx["l1"] != ty["l1"]:
        shared = 0
    elif tx["l2"] != ty["l2"]:
        shared = 1
    return shared / 3.0


sim_taste = np.zeros((n, n))
for i in range(n):
    for j in range(n):
        sim_taste[i, j] = taste_sim(taste[i], taste[j])

save_csv(sim_taste, "tax_taste_sim.csv", "sim")
heatmap(sim_taste, "tax_taste_sim.png", "Вкус — близость (У–Палмер, таксономия)")

# ==================================================================
# ЛР 2.3. Итоговая мера: взвешенная сумма частных близостей
# Несколько схем весов ω: экспертная + объективные альтернативы.
# ==================================================================
sim_num = 1.0 - d_euclid          # выбор ЛР 2.1 — норм. евклидово расстояние
sim_bin = binary_measures["sokal_michener"]   # выбор — SMC

PARTS = [
    ("taste",       sim_taste,       1),  # иерархия L1/L2/L3 — один признак
    ("numeric",     sim_num,         3),  # abv, sugar_gpl, volume_l
    ("binary",      sim_bin,         2),  # bitter, wood_contact
    ("sugar_class", sim_sugar_class, 1),
    ("country",     sim_country,     1),
    ("producer",    sim_producer,    1),
]
PART_KEYS = [p[0] for p in PARTS]
PART_MATS = [p[1] for p in PARTS]
PART_NFEAT = np.array([p[2] for p in PARTS], dtype=float)


def w_norm(v):
    v = np.asarray(v, dtype=float)
    s = v.sum()
    return v / s if s > 0 else np.ones_like(v) / len(v)


def shannon_H(vals, bins=10):
    """Энтропия Шеннона гистограммы значений верхнего треугольника."""
    hist, _ = np.histogram(vals, bins=bins, range=(0.0, 1.0))
    p = hist.astype(float)
    tot = p.sum()
    if tot == 0:
        return 0.0
    p = p / tot
    p = p[p > 0]
    return float(-(p * np.log(p)).sum())


def aggregate(w):
    S = np.zeros((n, n))
    for M, wk in zip(PART_MATS, w):
        S += wk * M
    return S


# --- схемы ω -------------------------------------------------------
# 1) экспертная — семантика рекомендации (вкус > профиль > флаги > …)
w_expert = np.array([0.35, 0.30, 0.15, 0.10, 0.05, 0.05])

# 2) равные — никакой априорной важности, ω_k = 1/K
w_equal = w_norm(np.ones(len(PARTS)))

# 3) по числу признаков в блоке (numeric 3, binary 2, остальные 1)
w_nfeat = w_norm(PART_NFEAT)

# 4) энтропия Шеннона распределения пар: разнообразнее шкала → больше вес
H_parts = np.array([shannon_H(upper_tri(M)) for M in PART_MATS])
w_entropy = w_norm(H_parts)

# 5) по СКО (контрастность): больший разброс пар → лучше разделяет
std_parts = np.array([upper_tri(M).std(ddof=1) for M in PART_MATS])
w_std = w_norm(std_parts)

# 6) CRITIC (Diakoulaki et al., 1995): контрастность × конфликт
#    C_k = σ_k · Σ_j (1 − ρ_kj); корреляции — по векторам верхних треугольников
vecs = np.column_stack([upper_tri(M) for M in PART_MATS])
R_parts = np.corrcoef(vecs, rowvar=False)
C_critic = np.array([
    std_parts[k] * np.sum(1.0 - R_parts[k]) for k in range(len(PARTS))
])
w_critic = w_norm(C_critic)

SCHEMES = {
    "expert":  ("экспертные",            w_expert),
    "equal":   ("равные",                w_equal),
    "nfeat":   ("по числу признаков",    w_nfeat),
    "entropy": ("энтропия Шеннона",      w_entropy),
    "std":     ("по СКО",                w_std),
    "critic":  ("CRITIC",                w_critic),
}

SIM_SCHEMES = {slug: aggregate(w) for slug, (_, w) in SCHEMES.items()}
sim_total = SIM_SCHEMES["expert"]
dist_total = 1.0 - sim_total

save_csv(sim_total, "agg_total_sim.csv", "sim")
save_csv(dist_total, "agg_total_dist.csv", "dist")
heatmap(sim_total, "agg_total_sim.png",
        "ИТОГО: близость (экспертные веса)")
heatmap(dist_total, "agg_total_dist.png",
        "ИТОГО: расстояние (1 - sim, экспертные веса)", cmap="magma")

for slug, (title, _w) in SCHEMES.items():
    if slug == "expert":
        continue
    S = SIM_SCHEMES[slug]
    save_csv(S, f"agg_{slug}_sim.csv", "sim")
    heatmap(S, f"agg_{slug}_sim.png",
            f"ИТОГО: близость (веса: {title})")

# Сводная тепловая карта всех схем (одна шкала [0,1], общая colorbar)
fig, axes = plt.subplots(2, 3, figsize=(15, 11))
im = None
for ax, (slug, (title, w)) in zip(axes.ravel(), SCHEMES.items()):
    im = ax.imshow(SIM_SCHEMES[slug], cmap="viridis", vmin=0.0, vmax=1.0)
    w1 = "  ".join(f"{k}={wk:.2f}" for k, wk in zip(PART_KEYS[:3], w[:3]))
    w2 = "  ".join(f"{k}={wk:.2f}" for k, wk in zip(PART_KEYS[3:], w[3:]))
    ax.set_title(f"{title}\n{w1}\n{w2}", fontsize=9)
    ax.set_xticks([])
    ax.set_yticks([])
fig.subplots_adjust(wspace=0.18, hspace=0.38, right=0.88, top=0.88)
cbar = fig.colorbar(im, ax=axes.ravel().tolist(), fraction=0.025, pad=0.04)
cbar.set_label("sim  [0, 1]")
fig.suptitle("Сравнение итоговых мер при разных схемах весов ω", fontsize=13)
fig.savefig(os.path.join(OUT_H, "agg_schemes_compare.png"), dpi=150,
            bbox_inches="tight")
plt.close(fig)

# |Δ| относительно экспертных — где схемы реально расходятся
fig, axes = plt.subplots(2, 3, figsize=(15, 11))
im = None
for ax, (slug, (title, _w)) in zip(axes.ravel(), SCHEMES.items()):
    D = np.abs(SIM_SCHEMES[slug] - sim_total)
    im = ax.imshow(D, cmap="magma", vmin=0.0, vmax=0.27)
    ax.set_title(title if slug != "expert" else "экспертные (|Δ|=0)", fontsize=11)
    ax.set_xticks([])
    ax.set_yticks([])
fig.subplots_adjust(wspace=0.18, hspace=0.28, right=0.88, top=0.90)
cbar = fig.colorbar(im, ax=axes.ravel().tolist(), fraction=0.025, pad=0.04)
cbar.set_label("|sim − sim_эксперт|")
fig.suptitle("Абсолютное расхождение с экспертными весами", fontsize=13)
fig.savefig(os.path.join(OUT_H, "agg_schemes_delta.png"), dpi=150,
            bbox_inches="tight")
plt.close(fig)

# Столбики весов
fig, ax = plt.subplots(figsize=(12, 5))
x = np.arange(len(PART_KEYS))
width = 0.13
for i, (slug, (title, w)) in enumerate(SCHEMES.items()):
    ax.bar(x + (i - 2.5) * width, w, width, label=title)
ax.set_xticks(x)
ax.set_xticklabels(PART_KEYS)
ax.set_ylabel("ω")
ax.set_ylim(0, 0.55)
ax.set_title("Веса частных мер при разных схемах")
ax.legend(fontsize=8, ncol=3)
ax.axhline(1.0 / len(PARTS), color="#888", ls="--", lw=0.8)
fig.tight_layout()
fig.savefig(os.path.join(OUT_H, "agg_weights_bar.png"), dpi=150)
plt.close(fig)


# ==================================================================
# Сводный xlsx: лист на каждую меру + цветовая шкала
# ==================================================================
def write_xlsx():
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, (M, kind) in ALL_MAT.items():
        ws = wb.create_sheet(name[:31])
        ws.append(["id"] + ids)
        for i in range(n):
            ws.append([labels[i]] + [round(float(M[i, j]), 4) for j in range(n)])
        rng = f"B2:{get_column_letter(n + 1)}{n + 1}"
        # sim: красный=далёкие ... зелёный=близкие; dist — наоборот
        lo, mid, hi = ("F8696B", "FFEB84", "63BE7B") if kind == "sim" \
            else ("63BE7B", "FFEB84", "F8696B")
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color=lo,
            mid_type="num", mid_value=0.5, mid_color=mid,
            end_type="num", end_value=1, end_color=hi))
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 26
        for c in range(2, n + 2):
            ws.column_dimensions[get_column_letter(c)].width = 6
    path = os.path.join(HERE, "out", "matrices.xlsx")
    wb.save(path)
    return path


# ==================================================================
# Интерактивная HTML-страница со всеми матрицами
# ==================================================================
HTML_TEMPLATE = """<!doctype html>
<html lang="ru"><head><meta charset="utf-8">
<title>Матрицы мер близости — ликёры</title>
<style>
body{font:13px/1.3 Consolas,monospace;margin:12px;background:#1e1e1e;color:#ddd}
select{font:inherit;margin:0 0 10px 0;padding:4px}
table{border-collapse:collapse}
td,th{padding:1px 4px;text-align:right;white-space:nowrap}
th.rot{vertical-align:bottom;height:150px;max-width:16px;width:16px}
th.rot>div{transform:rotate(-90deg) translateX(-4px);transform-origin:left bottom;
  width:150px;overflow:hidden;text-overflow:ellipsis}
td.rowh{text-align:left;position:sticky;left:0;background:#1e1e1e}
td.c{color:#111;min-width:44px}
td.c:hover{outline:2px solid #fff;cursor:default}
#tip{position:fixed;display:none;background:#333;border:1px solid #777;
  padding:6px 9px;pointer-events:none;z-index:9;max-width:340px}
#wrap{display:flex;align-items:flex-start;gap:14px}
#tbl-scroll{overflow:auto;max-width:calc(100vw - 90px)}
#cbar{position:sticky;top:12px;display:flex;flex-direction:column;align-items:center;
  flex-shrink:0;padding:4px 2px}
#cbar-title{font-size:11px;margin-bottom:8px;color:#bbb;letter-spacing:.04em}
#cbar-body{display:flex;align-items:stretch;gap:7px}
#cbar-grad{width:22px;height:320px;border:1px solid #666;border-radius:2px}
#cbar-ticks{display:flex;flex-direction:column;justify-content:space-between;
  height:320px;font-size:11px;color:#ccc}
#cbar-ticks span{line-height:1}
</style></head><body>
<select id="sel"></select>
<div id="cap"></div>
<div id="wrap">
  <div id="tbl-scroll"><table id="tbl"></table></div>
  <div id="cbar">
    <div id="cbar-title">sim</div>
    <div id="cbar-body">
      <div id="cbar-grad"></div>
      <div id="cbar-ticks">
        <span>1.0</span><span>0.8</span><span>0.6</span>
        <span>0.4</span><span>0.2</span><span>0.0</span>
      </div>
    </div>
  </div>
</div>
<div id="tip"></div>
<script>
const DATA = __DATA__;
const V = [[68,1,84],[59,82,139],[33,145,140],[94,201,98],[253,231,37]]; // viridis
const GRAD_SIM = 'linear-gradient(to top,rgb(68,1,84),rgb(59,82,139),rgb(33,145,140),rgb(94,201,98),rgb(253,231,37))';
const GRAD_DIST = 'linear-gradient(to top,rgb(253,231,37),rgb(94,201,98),rgb(33,145,140),rgb(59,82,139),rgb(68,1,84))';
function lerp(c){ return t=>{ t=Math.min(Math.max(t,0),1)*(c.length-1);
  const i=Math.min(Math.floor(t),c.length-2), f=t-i;
  return c[i].map((v,k)=>Math.round(v+(c[i+1][k]-v)*f)); }; }
const colSim = lerp(V), colDist = lerp([...V].reverse());
const sel=document.getElementById('sel'), tbl=document.getElementById('tbl'),
      tip=document.getElementById('tip'), cap=document.getElementById('cap'),
      cbarGrad=document.getElementById('cbar-grad'),
      cbarTitle=document.getElementById('cbar-title');
for(const k of Object.keys(DATA.matrices)){
  const o=document.createElement('option'); o.value=k; o.textContent=k; sel.appendChild(o);}
function render(name){
  const {kind,data}=DATA.matrices[name], cf=kind==='sim'?colSim:colDist;
  cap.textContent=name+' — '+(kind==='sim'?'близость':'расстояние')+
    ' · жёлтый — близкие, фиолетовый — далёкие';
  cbarTitle.textContent = kind==='sim' ? 'sim' : 'dist';
  cbarGrad.style.background = kind==='sim' ? GRAD_SIM : GRAD_DIST;
  let h='<tr><th></th>'+DATA.ids.map((id,j)=>'<th class="rot"><div>'+id+'</div></th>').join('')+'</tr>';
  for(let i=0;i<DATA.ids.length;i++){
    h+='<tr><td class="rowh">'+DATA.ids[i]+' '+DATA.names[i]+'</td>';
    for(let j=0;j<DATA.ids.length;j++){
      const v=data[i][j], rgb=cf(v);
      h+='<td class="c" data-i="'+i+'" data-j="'+j+'" data-v="'+v+'" '+
         'style="background:rgb('+rgb+')">'+v.toFixed(2)+'</td>';
    }
    h+='</tr>';
  }
  tbl.innerHTML=h;
}
tbl.addEventListener('mousemove',e=>{
  const c=e.target.closest('td.c'); if(!c){tip.style.display='none';return;}
  const i=+c.dataset.i,j=+c.dataset.j;
  tip.innerHTML='<b>'+DATA.ids[i]+'</b> '+DATA.names[i]+'<br><b>'+DATA.ids[j]+'</b> '+DATA.names[j]+'<br>значение: <b>'+c.dataset.v+'</b>';
  tip.style.display='block'; tip.style.left=(e.clientX+14)+'px'; tip.style.top=(e.clientY+14)+'px';
});
tbl.addEventListener('mouseleave',()=>tip.style.display='none');
sel.onchange=()=>{render(sel.value);tip.style.display='none';};
render(sel.value=Object.keys(DATA.matrices)[0]);
</script></body></html>"""


def write_html():
    payload = {
        "ids": ids,
        "names": names,
        "matrices": {k: {"kind": kind, "data": np.round(M, 3).tolist()}
                     for k, (M, kind) in ALL_MAT.items()},
    }
    html = HTML_TEMPLATE.replace("__DATA__", json.dumps(payload, ensure_ascii=False))
    path = os.path.join(HERE, "out", "matrices.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


print("xlsx:", write_xlsx())
print("html:", write_html())


# ==================================================================
# Экспертная оценка (числа для отчёта)
# ==================================================================
idx = {rid: i for i, rid in enumerate(ids)}

ref_pairs = [
    ("L26", "L27", "Crème de Cacao Brown/White: числовые признаки идентичны"),
    ("L08", "L10", "Cointreau/Grand Marnier: 40%, сахар 200 vs 214"),
    ("L20", "L21", "Disaronno/Luxardo: 28%, сахар 300 vs 280"),
    ("L15", "L16", "Абрикос/персик Giffard/Archers: сахар 280, abv 25 vs 18"),
    ("L04", "L13", "Chartreuse 55% vs Ponche 15%: крайние по крепости"),
    ("L06", "L30", "Fernet 39%/80 сахар vs Advocaat 17.2%/260"),
    ("L01", "L19", "Мята/банан: сахар оба 400, abv 24 vs 17"),
]

rep("=" * 78)
rep("ЛР 2.1a. СРАВНЕНИЕ ЧИСЛОВЫХ МЕР НА ЭТАЛОННЫХ ПАРАХ")
rep("=" * 78)
header = (f'{"пара":12} {"евклид":>8} {"манхэтт":>8} {"минк3":>8} '
          f'{"чебышёв":>8} {"косинус":>8}   комментарий')
rep(header)
for i_id, j_id, note in ref_pairs:
    i, j = idx[i_id], idx[j_id]
    rep(f"{i_id}-{j_id:5}  {d_euclid[i,j]:8.4f} {d_manh[i,j]:8.4f} "
        f"{d_mink3[i,j]:8.4f} {d_cheb[i,j]:8.4f} {d_cosine[i,j]:8.4f}   {note}")

rep()
rep("Спирменовы корреляции ранжирований пар между мерами:")
ms = list(numeric_measures)
for a_i in range(len(ms)):
    for b_i in range(a_i + 1, len(ms)):
        rho = spearman(upper_tri(numeric_measures[ms[a_i]]),
                       upper_tri(numeric_measures[ms[b_i]]))
        rep(f"  {ms[a_i]:>10} vs {ms[b_i]:<10} rho = {rho:.4f}")

rep()
rep("Доля пар с одинаковым расстоянием (связки, грубость шкалы):")
for name_m, M in numeric_measures.items():
    vals = np.round(upper_tri(M), 6)
    uniq, cnt = np.unique(vals, return_counts=True)
    tied = cnt[cnt > 1].sum()
    rep(f"  {name_m:>10}: уникальных значений {len(uniq):3d} из {len(vals)}, "
        f"значений со связями {tied}")

rep()
rep("Ближайшие соседи (top-3) по каждой числовой мере для ключевых ликёров:")
probe = ["L26", "L05", "L28", "L24", "L08"]
for rid in probe:
    i = idx[rid]
    rep(f"  {rid} {names[i]}")
    for name_m, M in numeric_measures.items():
        order = np.argsort(M[i] + np.eye(n)[i] * 1e9)[:3]
        rep(f"    {name_m:>10}: " +
            ", ".join(f"{ids[k]}({M[i,k]:.3f})" for k in order))

rep()
rep("=" * 78)
rep("ЛР 2.1b. АССОЦИАТИВНЫЕ БИНАРНЫЕ МЕРЫ — сравнение")
rep("=" * 78)
rep("Пары, где оба объекта (0,0) по bitter/wood — тест на полноту учёта:")
zero_rows = [i for i in range(n) if B[i].sum() == 0]
rep(f"  объектов с (0,0): {len(zero_rows)}")
i, j = idx["L26"], idx["L27"]
for name_m, M in binary_measures.items():
    rep(f"  L26-L27 {name_m:>15}: {M[i,j]:.3f}")
rep("Пара с одним рассогласованием (L28(0,1) vs L30(0,0)):")
i, j = idx["L28"], idx["L30"]
for name_m, M in binary_measures.items():
    rep(f"  L28-L30 {name_m:>15}: {M[i,j]:.3f}")
rep("Пара с полным рассогласованием (L05(1,1) vs L01(0,0)):")
i, j = idx["L05"], idx["L01"]
for name_m, M in binary_measures.items():
    rep(f"  L05-L01 {name_m:>15}: {M[i,j]:.3f}")

rep()
rep("=" * 78)
rep("ЛР 2.2. ТАКСОНОМИЯ — примеры значений")
rep("=" * 78)
tax_pairs = [
    ("L26", "L27", "общий лист (R-C-*)"),
    ("L20", "L21", "общий L2 'Ядра косточковых', разный L3"),
    ("L05", "L07", "общий лист 'Амаро'"),
    ("L01", "L03", "общий L1 'Травяные', разный L2"),
    ("L08", "L26", "разный L1 (Фруктовые vs Обжаренные)"),
]
for i_id, j_id, note in tax_pairs:
    i, j = idx[i_id], idx[j_id]
    rep(f"  {i_id}-{j_id}: sim={sim_taste[i,j]:.3f}   {note}")

rep()
rep("=" * 78)
rep("ЛР 2.3. АГРЕГИРОВАННАЯ МЕРА — top-5 соседей (экспертные веса)")
rep("=" * 78)
for rid in ["L26", "L05", "L28", "L03", "L08", "L24"]:
    i = idx[rid]
    order = np.argsort(-(sim_total[i] - np.eye(n)[i] * 2))[:5]
    rep(f"  {rid} {names[i]}:")
    rep("      " + ", ".join(f"{ids[k]} {names[k][:20]} ({sim_total[i,k]:.3f})"
                            for k in order))

rep()
rep("=" * 78)
rep("ЛР 2.3. СХЕМЫ ВЕСОВ ω — сравнение")
rep("=" * 78)
header = f'{"схема":22}' + "".join(f"{k:>12}" for k in PART_KEYS)
rep(header)
for slug, (title, w) in SCHEMES.items():
    rep(f"{title:22}" + "".join(f"{wk:12.3f}" for wk in w))

rep()
rep("Диагностика частных мер (верхний треугольник, 435 пар):")
rep(f'  {"мера":12} {"H":>8} {"σ":>8} {"CRITIC C":>10}  уникальных')
for k, key in enumerate(PART_KEYS):
    vals = np.round(upper_tri(PART_MATS[k]), 6)
    uniq = len(np.unique(vals))
    rep(f"  {key:12} {H_parts[k]:8.3f} {std_parts[k]:8.3f} "
        f"{C_critic[k]:10.3f}  {uniq:3d}")

rep()
rep("Корреляции Пирсона частных мер (верхний треугольник):")
rep("         " + "".join(f"{k:>12}" for k in PART_KEYS))
for i, ki in enumerate(PART_KEYS):
    rep(f"  {ki:7}" + "".join(f"{R_parts[i, j]:12.3f}" for j in range(len(PART_KEYS))))

rep()
rep("Спирмен ρ итоговых матриц (ранжирование 435 пар vs экспертные):")
expert_vec = upper_tri(sim_total)
for slug, (title, _w) in SCHEMES.items():
    if slug == "expert":
        continue
    rho = spearman(expert_vec, upper_tri(SIM_SCHEMES[slug]))
    off = np.abs(sim_total - SIM_SCHEMES[slug])
    iu = np.triu_indices(n, 1)
    rep(f"  {title:22}  ρ = {rho:.4f}   "
        f"mean|Δ|={off[iu].mean():.4f}  max|Δ|={off[iu].max():.4f}")

rep()
rep("Спирмен ρ между всеми схемами:")
slugs = list(SCHEMES)
rep("         " + "".join(f"{s:>10}" for s in slugs))
for a in slugs:
    row = []
    for b in slugs:
        row.append(spearman(upper_tri(SIM_SCHEMES[a]), upper_tri(SIM_SCHEMES[b])))
    rep(f"  {a:7}" + "".join(f"{v:10.3f}" for v in row))

rep()
rep("Top-1 сосед vs экспертные (ключевые объекты):")
for rid in ["L26", "L05", "L28", "L03", "L08", "L24"]:
    i = idx[rid]
    exp_nn = np.argsort(-(sim_total[i] - np.eye(n)[i] * 2))[0]
    bits = [f"эксперт={ids[exp_nn]}"]
    for slug, (title, _w) in SCHEMES.items():
        if slug == "expert":
            continue
        S = SIM_SCHEMES[slug]
        nn = np.argsort(-(S[i] - np.eye(n)[i] * 2))[0]
        mark = "" if nn == exp_nn else " *"
        bits.append(f"{title[:8]}={ids[nn]}{mark}")
    rep(f"  {rid}: " + ", ".join(bits))

rep()
rep("Сколько объектов из 30 меняют ближайшего соседа относительно экспертных:")
for slug, (title, _w) in SCHEMES.items():
    if slug == "expert":
        continue
    changed, examples = 0, []
    for i in range(n):
        exp_nn = np.argsort(-(sim_total[i] - np.eye(n)[i] * 2))[0]
        nn = np.argsort(-(SIM_SCHEMES[slug][i] - np.eye(n)[i] * 2))[0]
        if nn != exp_nn:
            changed += 1
            if len(examples) < 4:
                examples.append(f"{ids[i]}:{ids[exp_nn]}→{ids[nn]}")
    extra = ("  " + ", ".join(examples)) if examples else "  (никто)"
    rep(f"  {title:22}  {changed:2d}/30{extra}")

rep()
rep("Пары с наибольшим |Δ| относительно экспертных:")
for slug, (title, _w) in SCHEMES.items():
    if slug == "expert":
        continue
    D = np.abs(sim_total - SIM_SCHEMES[slug])
    np.fill_diagonal(D, 0.0)
    k = int(np.argmax(D))
    i, j = divmod(k, n)
    rep(f"  {title:22}  {ids[i]}-{ids[j]}  "
        f"эксперт={sim_total[i, j]:.3f}  схема={SIM_SCHEMES[slug][i, j]:.3f}  "
        f"Δ={SIM_SCHEMES[slug][i, j] - sim_total[i, j]:+.3f}")

rep()
rep("Проверки: diag(sim)=1, симметрия, диапазоны")
for nm, M in [("sim_num", sim_num), ("sim_bin", sim_bin),
              ("sim_taste", sim_taste), ("sim_total", sim_total)]:
    rep(f"  {nm}: diag=[{np.diag(M).min():.3f},{np.diag(M).max():.3f}] "
        f"sym={np.allclose(M, M.T)} range=[{M.min():.3f},{M.max():.3f}]")
for slug, (title, _w) in SCHEMES.items():
    M = SIM_SCHEMES[slug]
    rep(f"  agg/{slug}: diag=[{np.diag(M).min():.3f},{np.diag(M).max():.3f}] "
        f"sym={np.allclose(M, M.T)} range=[{M.min():.3f},{M.max():.3f}]")

with open(OUT_R, "w", encoding="utf-8") as f:
    f.write("\n".join(report))

print(f"\nГотово. CSV -> {OUT_M}\nPNG  -> {OUT_H}\nОтчёт -> {OUT_R}")
